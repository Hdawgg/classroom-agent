"""
Read-only Google Classroom + Drive reader for study mode.

Loads every course with its coursework, materials and announcements, and extracts the text of attached
Drive files (Google Docs/Slides, PDF, .docx, .pptx, plain text). Extracted text is cached in .study_cache/
so repeat runs don't re-download unchanged files. Anything unreadable is skipped with the reason recorded.
"""

import io
import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, List

import google_classroom
from google_classroom import BASE_DIR, CREDS_FILE, TOKEN_FILE

CACHE_DIR = BASE_DIR / ".study_cache"
SAMPLE_FILE = BASE_DIR / "study_sample.json"
MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
# pypdf logs warnings about odd PDFs to the console; failures are already recorded as skip reasons
logging.getLogger("pypdf").setLevel(logging.ERROR)

STUDY_SCOPES = [
    "https://www.googleapis.com/auth/classroom.courses.readonly",
    "https://www.googleapis.com/auth/classroom.coursework.me.readonly",
    "https://www.googleapis.com/auth/classroom.courseworkmaterials.readonly",
    "https://www.googleapis.com/auth/classroom.announcements.readonly",
    # courseWork only carries a topicId; this read-only scope turns it into the topic's name
    "https://www.googleapis.com/auth/classroom.topics.readonly",
    "https://www.googleapis.com/auth/drive.readonly",
]
# A fresh login asks for both modes' scopes so one token.json serves solve mode and study mode
LOGIN_SCOPES = list(dict.fromkeys(google_classroom.SCOPES + STUDY_SCOPES))

GDOC = "application/vnd.google-apps.document"
GSLIDES = "application/vnd.google-apps.presentation"
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
KIND_BY_EXT = {".pdf": PDF, ".docx": DOCX, ".pptx": PPTX}
TEXT_EXTS = {".txt", ".md", ".py", ".c", ".cpp", ".h", ".java", ".js", ".csv", ".json", ".sql", ".html"}
FRIENDLY_MIME = {
    GDOC: "Google Doc", GSLIDES: "Google Slides", PDF: "PDF", DOCX: "Word .docx", PPTX: "PowerPoint .pptx",
    "application/vnd.google-apps.spreadsheet": "Google Sheets", "application/vnd.google-apps.form": "Google Form",
    "application/vnd.google-apps.folder": "Drive folder", "application/msword": "old Word .doc",
    "application/vnd.ms-powerpoint": "old PowerPoint .ppt",
}


class ReauthNeeded(RuntimeError):
    """token.json/credentials.json can't give study mode the access it needs; the message says what to do."""


def _short(scope: str) -> str:
    return scope.rsplit("/", 1)[-1]


def missing_scopes(token_file: Path = TOKEN_FILE) -> List[str]:
    """Study scopes that token.json was not granted (all of them if the file is unreadable)."""
    try:
        granted = json.loads(token_file.read_text(encoding="utf-8")).get("scopes") or []
    except (OSError, ValueError, AttributeError):
        return list(STUDY_SCOPES)
    if isinstance(granted, str):
        granted = granted.split()
    granted = set(granted)
    # A read-write scope (e.g. classroom.coursework.me) also covers its .readonly variant
    return [s for s in STUDY_SCOPES if s not in granted and s.removesuffix(".readonly") not in granted]


def login():
    """Returns (classroom, drive) API clients, or raises ReauthNeeded with instructions."""
    if TOKEN_FILE.exists():
        missing = missing_scopes()
        if missing:
            raise ReauthNeeded(
                f"{TOKEN_FILE.name} was created without the read-only access study mode needs:\n"
                + "".join(f"  - {_short(s)}\n" for s in missing)
                + f"Delete {TOKEN_FILE.name} and run study mode again to sign in. On Google's consent screen, tick every box.\n"
                "The new sign-in also covers solve mode."
            )
    elif not CREDS_FILE.exists():
        raise ReauthNeeded("credentials.json is missing (see 'Google Classroom Setup' in README.md). "
                           "Use --offline to try study mode on the sample data.")

    # scopes=None keeps the scopes token.json was granted, so a refresh never asks for more than that
    creds = google_classroom.get_credentials(scopes=None, flow_scopes=LOGIN_SCOPES)
    granted = getattr(creds, "granted_scopes", None)
    if granted:  # only known right after a fresh browser sign-in
        unticked = [s for s in STUDY_SCOPES if s not in set(granted)]
        if unticked:
            print("[WARN] These permissions were not ticked on the consent screen: "
                  + ", ".join(_short(s) for s in unticked)
                  + f". Items that need them will be skipped. To fix: delete {TOKEN_FILE.name} and sign in again.")

    from googleapiclient.discovery import build
    return (build("classroom", "v1", credentials=creds, cache_discovery=False),
            build("drive", "v3", credentials=creds, cache_discovery=False))


# ============================================================
# CLASSROOM: courses and their items
# ============================================================
def _http_reason(e: Exception) -> str:
    status = getattr(getattr(e, "resp", None), "status", None)
    reason = getattr(e, "reason", "") or str(e)
    return f"HTTP {status}: {reason}" if status else reason


def fetch_courses(classroom, course_filter: str = "all") -> List[Dict[str, Any]]:
    courses = google_classroom._list_all(classroom.courses().list, "courses", studentId="me", courseStates=["ACTIVE"])
    return [
        {"id": c["id"], "name": c.get("name", "Unknown Course"), "section": c.get("section", ""), "items": None, "warnings": []}
        for c in courses if google_classroom.course_matches(c.get("name", ""), course_filter)
    ]


def _attachments(materials: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out = []
    for m in materials or []:
        if "driveFile" in m:
            f = m["driveFile"].get("driveFile", {})
            out.append({"type": "drive", "file_id": f.get("id"), "title": f.get("title", "Drive file"), "url": f.get("alternateLink", "")})
        elif "youtubeVideo" in m:
            v = m["youtubeVideo"]
            out.append({"type": "youtube", "title": v.get("title", "YouTube video"), "url": v.get("alternateLink", "")})
        elif "link" in m:
            l = m["link"]
            out.append({"type": "link", "title": l.get("title") or l.get("url", "Link"), "url": l.get("url", "")})
        elif "form" in m:
            f = m["form"]
            out.append({"type": "form", "title": f.get("title", "Google Form"), "url": f.get("formUrl", "")})
    for a in out:
        if a["type"] != "drive":
            a.update(status="not_fetched", reason=f"{a['type']} not fetched (only Drive files are read)")
    return out


def fetch_course_items(classroom, course: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Coursework, materials and announcements of one course (attachments not read yet).
    A list the API refuses (e.g. a missing permission) is recorded in course['warnings'] and skipped."""
    c_id, c_name = course["id"], course["name"]
    courses_api = classroom.courses()

    def safe_list(label, request_fn, key, **kwargs):
        try:
            return google_classroom._list_all(request_fn, key, courseId=c_id, **kwargs)
        except Exception as e:
            course["warnings"].append(f"Could not load {label}: {_http_reason(e)}")
            return []

    topics = {t["topicId"]: t.get("name", "") for t in safe_list("topics", courses_api.topics().list, "topic")}
    items = []
    for kind, request_fn, key in [
        ("Assignment", courses_api.courseWork().list, "courseWork"),
        ("Material", courses_api.courseWorkMaterials().list, "courseWorkMaterial"),
        ("Announcement", courses_api.announcements().list, "announcements"),
    ]:
        for raw in safe_list(f"{kind.lower()}s", request_fn, key):
            text = (raw.get("description") if kind != "Announcement" else raw.get("text")) or ""
            # Announcements have no title; study.py builds a display label from their text and files
            title = (raw.get("title") or "").strip()
            items.append({
                "id": raw["id"], "course_id": c_id, "course": c_name, "kind": kind, "title": title,
                "description": text.strip(),
                "due": google_classroom._format_due_date(raw.get("dueDate")) if raw.get("dueDate") else "",
                "topic": topics.get(raw.get("topicId"), "(Unknown topic)" if raw.get("topicId") else ""),
                "created": raw.get("creationTime", ""),
                "link": raw.get("alternateLink", ""),
                "attachments": _attachments(raw.get("materials")),
            })
    return items


# ============================================================
# DRIVE: download/export + text extraction (cached)
# ============================================================
def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "").lstrip("﻿")  # Docs exports start with a BOM
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("password-protected PDF")
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for n, slide in enumerate(prs.slides, 1):
        parts.append(f"\n--- Slide {n} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
    return "\n".join(parts)


EXTRACTORS = {PDF: extract_pdf, DOCX: extract_docx, PPTX: extract_pptx}


def _cache_file(file_id: str) -> Path:
    return CACHE_DIR / "drive" / (re.sub(r"[^A-Za-z0-9_-]", "_", file_id) + ".json")


def read_drive_file(drive, att: Dict[str, Any], refresh: bool = False) -> None:
    """Fills att with status ('read'/'skipped'), text or reason, and mime. Never raises."""
    file_id = att.get("file_id")
    if not file_id:
        att.update(status="skipped", reason="attachment has no Drive file ID")
        return
    try:
        meta = drive.files().get(fileId=file_id, supportsAllDrives=True,
                                 fields="id,name,mimeType,modifiedTime,size,capabilities/canDownload").execute()
    except Exception as e:
        att.update(status="skipped", reason=f"could not open file ({_http_reason(e)})")
        return

    mime, name = meta.get("mimeType", ""), meta.get("name", att["title"])
    att["mime"] = FRIENDLY_MIME.get(mime, mime)
    if meta.get("size"):  # Google Docs/Slides have no size
        att["size"] = int(meta["size"])
    cache_file = _cache_file(file_id)
    if not refresh and cache_file.exists():
        try:
            cached = json.loads(cache_file.read_text(encoding="utf-8"))
            if cached.get("modifiedTime") == meta.get("modifiedTime"):
                att.update(status=cached["status"], text=cached.get("text", ""), reason=cached.get("reason", ""), cached=True)
                return
        except (OSError, ValueError, KeyError):
            pass  # corrupt cache entry: re-download

    ext = Path(name).suffix.lower()
    kind = KIND_BY_EXT.get(ext, mime) if mime in ("application/octet-stream", "") else mime
    try:
        if kind in (GDOC, GSLIDES):
            data = drive.files().export(fileId=file_id, mimeType="text/plain").execute()
            text = data.decode("utf-8", errors="replace") if isinstance(data, bytes) else str(data)
        elif kind in EXTRACTORS or kind.startswith("text/") or ext in TEXT_EXTS:
            if meta.get("capabilities", {}).get("canDownload") is False:
                att.update(status="skipped", reason="the owner disabled downloading for this file")
                return
            if int(meta.get("size") or 0) > MAX_DOWNLOAD_BYTES:
                att.update(status="skipped", reason=f"file is larger than {MAX_DOWNLOAD_BYTES // (1024 * 1024)} MB")
                return
            data = drive.files().get_media(fileId=file_id, supportsAllDrives=True).execute()
            text = EXTRACTORS[kind](data) if kind in EXTRACTORS else data.decode("utf-8", errors="replace")
        else:
            att.update(status="skipped", reason=f"unsupported file type ({att['mime']})")
            return
    except Exception as e:
        if hasattr(e, "resp"):  # HTTP errors may be temporary, so they are not cached
            att.update(status="skipped", reason=f"download failed ({_http_reason(e)})")
            return
        result = {"status": "skipped", "reason": f"could not extract text ({type(e).__name__}: {e})"}
    else:
        text = _clean(text)
        if text:
            result = {"status": "read", "text": text}
        else:
            hint = " - probably a scanned PDF (images only)" if kind == PDF else ""
            result = {"status": "skipped", "reason": f"no extractable text{hint}"}

    att.update(result)
    # Results for this exact file version are cached, including permanent failures, so they aren't retried every run
    try:
        cache_file.parent.mkdir(parents=True, exist_ok=True)
        cache_file.write_text(json.dumps({"name": name, "mime": mime, "modifiedTime": meta.get("modifiedTime"), **result}),
                              encoding="utf-8")
    except OSError:
        pass


def read_attachments(drive, items: List[Dict[str, Any]], refresh: bool = False, progress=print) -> None:
    drive_atts = [a for item in items for a in item["attachments"] if a["type"] == "drive" and "status" not in a]
    for n, att in enumerate(drive_atts, 1):
        progress(f"  [{n}/{len(drive_atts)}] {att['title'][:60]}")
        read_drive_file(drive, att, refresh=refresh)


# ============================================================
# OFFLINE SAMPLE
# ============================================================
def load_sample() -> List[Dict[str, Any]]:
    """Courses with items and already-'read' attachments, in the same shape as live data."""
    data = json.loads(SAMPLE_FILE.read_text(encoding="utf-8"))
    courses = data.get("courses", [])
    for course in courses:
        course.setdefault("warnings", [])
        for item in course.get("items", []):
            item.setdefault("course_id", course["id"])
            item.setdefault("course", course["name"])
            for key in ("title", "description", "due", "topic", "created", "link"):
                item.setdefault(key, "")
            item.setdefault("attachments", [])
    return courses
